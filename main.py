# main.py —— 企业销售报表自动化工具
"""
读取 销售订单.xlsx
  → 数据体检（缺失值 / 重复 / 金额不符 / 日期异常）
  → 多维度汇总
  → 生成图表
  → 输出 销售报表.xlsx
"""
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).parent
SRC_FILE = BASE_DIR / "销售订单.xlsx"      # ← 换成你的真实文件名
OUT_FILE = BASE_DIR / "销售报表.xlsx"

REQUIRED = [
    "订单号", "下单日期", "客户名称", "区域",
    "产品编码", "产品名称", "单价", "数量", "金额", "销售员",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")


# ---------- 1. 读取 ----------
def read_source(path):
    if not path.exists():
        raise FileNotFoundError(f"找不到数据文件：{path}")

    raw = pd.read_excel(path)
    raw.columns = [str(c).strip() for c in raw.columns]

    missing = [c for c in REQUIRED if c not in raw.columns]
    if missing:
        raise ValueError(f"表里缺少这些列：{missing}")
    return raw


# ---------- 2. 数据体检 ----------
def check_quality(raw):
    """逐项体检，返回（检查汇总表, 问题明细表）"""
    checks = []
    flags = {}

    def flag(mask, name, note):
        mask = mask.fillna(False).astype(bool)
        rows = raw.index[mask]
        checks.append({"检查项": name, "问题行数": len(rows), "说明": note})
        for i in rows:
            flags.setdefault(i, []).append(name)

    for col in raw.columns:
        flag(raw[col].isna(), f"缺失值：{col}", f"「{col}」为空")

    order = raw["订单号"].astype("string").str.strip()
    flag(order.isna() | order.eq(""), "订单号为空", "可能是表尾合计行或空行")
    flag(order.duplicated(keep=False) & order.notna() & order.ne(""),
         "订单号重复", "同一订单号出现多次（可能是多行明细，需人工确认）")
    flag(raw.duplicated(), "整行重复", "整行数据完全重复，已自动去重")

    dates = pd.to_datetime(raw["下单日期"], errors="coerce")
    flag(dates.isna() & raw["下单日期"].notna(), "日期无法解析", "不是有效日期")

    price = pd.to_numeric(raw["单价"], errors="coerce")
    qty = pd.to_numeric(raw["数量"], errors="coerce")
    amount = pd.to_numeric(raw["金额"], errors="coerce")

    flag(price.isna() & raw["单价"].notna(), "单价无法解析", "单价含非数字")
    flag(qty.isna() & raw["数量"].notna(), "数量无法解析", "数量含非数字")
    flag(amount.isna() & raw["金额"].notna(), "金额无法解析", "金额含非数字")
    flag(price <= 0, "单价异常", "单价 ≤ 0")
    flag(qty <= 0, "数量异常", "数量 ≤ 0")
    flag((price * qty - amount).abs() > 0.01, "金额不符", "金额 ≠ 单价 × 数量")

    summary = pd.DataFrame(checks, columns=["检查项", "问题行数", "说明"])

    if flags:
        idx = sorted(flags)
        detail = raw.loc[idx].copy()
        detail.insert(0, "问题说明", ["；".join(flags[i]) for i in idx])
        detail.insert(0, "原表行号", [i + 2 for i in idx])
    else:
        detail = pd.DataFrame(columns=["原表行号", "问题说明", *raw.columns])

    return summary, detail


# ---------- 3. 清洗 ----------
def clean(raw):
    df = raw.copy()
    df["下单日期"] = pd.to_datetime(df["下单日期"], errors="coerce")
    for col in ("单价", "数量", "金额"):
        df[col] = pd.to_numeric(df[col], errors="coerce")

    before = len(df)
    df = df.drop_duplicates()
    df = df.dropna(subset=["下单日期", "金额", "订单号", "数量"])
    df = df[df["金额"] > 0]
    df = df[df["数量"] > 0]
    df = df[df["订单号"].astype(str).str.strip().ne("")]

    print(f"清洗：{before} 行 → {len(df)} 行（移除 {before - len(df)} 行）")

    df = df.copy()
    df["月份"] = df["下单日期"].dt.strftime("%Y-%m")
    return df


# ---------- 4. 汇总 ----------
def summarize(df, key, top=None):
    g = (
        df.groupby(key, dropna=False)
        .agg(销售额=("金额", "sum"),
             订单数=("订单号", "nunique"),
             销量=("数量", "sum"))
        .sort_values("销售额", ascending=False)
        .reset_index()
    )
    g["占比"] = g["销售额"] / g["销售额"].sum()
    g["客单价"] = g["销售额"] / g["订单数"]
    if top:
        g = g.head(top)
    return g.round(2)


def monthly(df):
    g = (
        df.groupby("月份")
        .agg(销售额=("金额", "sum"),
             订单数=("订单号", "nunique"),
             销量=("数量", "sum"))
        .sort_index()
        .reset_index()
    )
    g["环比"] = g["销售额"].pct_change()
    g["占比"] = g["销售额"] / g["销售额"].sum()
    return g.round(2)


def add_total(g, key, df):
    total = {
        key: "合计",
        "销售额": round(df["金额"].sum(), 2),
        "订单数": int(df["订单号"].nunique()),
        "销量": int(df["数量"].sum()),
        "占比": 1.0,
        "客单价": round(df["金额"].sum() / df["订单号"].nunique(), 2),
    }
    total = {col: total.get(col) for col in g.columns}
    return pd.concat([g, pd.DataFrame([total])], ignore_index=True)


def overview(df):
    total = df["金额"].sum()
    orders = int(df["订单号"].nunique())
    rows = [
        ("总销售额", round(total, 2)),
        ("订单数", orders),
        ("总销量", int(df["数量"].sum())),
        ("客户数", int(df["客户名称"].nunique())),
        ("产品数", int(df["产品名称"].nunique())),
        ("客单价", round(total / orders, 2)),
        ("平均每单数量", round(df["数量"].sum() / orders, 2)),
        ("数据区间", f'{df["下单日期"].min():%Y-%m-%d} 至 {df["下单日期"].max():%Y-%m-%d}'),
    ]
    return pd.DataFrame(rows, columns=["指标", "数值"])


# ---------- 5. 美化 ----------
def text_width(value):
    if value is None:
        return 0
    return sum(2 if ord(ch) > 127 else 1 for ch in str(value))


def style_sheet(ws, warn_positive=False):
    headers = [c.value for c in ws[1]]

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"

    for idx, cells in enumerate(ws.columns, start=1):
        width = max(text_width(c.value) for c in cells)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 4, 10), 36)

    for row in ws.iter_rows(min_row=2):
        is_total = str(row[0].value).strip() == "合计"
        warn = (
            warn_positive
            and len(row) > 1
            and isinstance(row[1].value, (int, float))
            and row[1].value > 0
        )
        for cell, header in zip(row, headers):
            if warn:
                cell.fill = WARN_FILL
            if is_total:
                cell.font = Font(bold=True)
                cell.fill = TOTAL_FILL
            if isinstance(cell.value, (int, float)):
                if header in ("占比", "环比"):
                    cell.number_format = "0.0%"
                elif "额" in str(header) or "价" in str(header):
                    cell.number_format = "#,##0.00"
                else:
                    cell.number_format = "#,##0"


def style_overview(ws):
    style_sheet(ws)
    for row in ws.iter_rows(min_row=2, max_col=2):
        if not isinstance(row[1].value, (int, float)):
            continue
        label = str(row[0].value)
        row[1].number_format = "#,##0.00" if ("额" in label or "价" in label) else "#,##0"


# ---------- 6. 图表 ----------
def add_bar(ws, title, n_rows, anchor="J2", cat_col=1, val_col=2):
    """柱状图：第 1 行是表头，数据从第 2 行开始，共 n_rows 行"""
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = title
    chart.y_axis.title = "销售额（元）"

    data = Reference(ws, min_col=val_col, min_row=1, max_row=1 + n_rows)
    cats = Reference(ws, min_col=cat_col, min_row=2, max_row=1 + n_rows)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    chart.width = 18
    chart.height = 9
    ws.add_chart(chart, anchor)


def add_pie(ws, title, n_rows, anchor="J2", cat_col=1, val_col=2):
    chart = PieChart()
    chart.title = title

    data = Reference(ws, min_col=val_col, min_row=1, max_row=1 + n_rows)
    cats = Reference(ws, min_col=cat_col, min_row=2, max_row=1 + n_rows)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    labels = DataLabelList()
    labels.showPercent = True
    chart.dataLabels = labels

    chart.width = 14
    chart.height = 9
    ws.add_chart(chart, anchor)


# ---------- 7. 主流程 ----------
def main():
    print("企业销售报表自动化工具")
    print("-" * 44)

    raw = read_source(SRC_FILE)
    print(f"读取 {SRC_FILE.name}：{len(raw)} 行 × {len(raw.columns)} 列")

    quality, problems = check_quality(raw)
    bad = quality[quality["问题行数"] > 0]

    print(f"\n数据体检：共 {len(quality)} 项检查，{len(problems)} 行存在问题")
    if len(bad):
        for _, r in bad.iterrows():
            print(f"  · {r['检查项']}：{r['问题行数']} 行（{r['说明']}）")
    else:
        print("  · 未发现数据问题")

    df = clean(raw)

    tables = {
        "总览": overview(df),
        "数据质量": quality,
        "按月份": add_total(monthly(df), "月份", df),
        "按区域": add_total(summarize(df, "区域"), "区域", df),
        "按产品": add_total(summarize(df, "产品名称"), "产品名称", df),
        "按销售员": add_total(summarize(df, "销售员"), "销售员", df),
        "客户TOP10": summarize(df, "客户名称", top=10),
        "明细": df.drop(columns=["月份"]).sort_values("下单日期", ascending=False),
    }
    if len(problems):
        tables["问题数据"] = problems

    with pd.ExcelWriter(OUT_FILE, engine="openpyxl") as writer:
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name, index=False)

        style_overview(writer.sheets["总览"])
        for name in tables:
            if name != "总览":
                style_sheet(writer.sheets[name], warn_positive=(name == "数据质量"))

        # 图表：数据行数 = 表格行数 - 合计行（不把合计画进图里）
        add_bar(writer.sheets["按月份"], "月度销售额趋势", len(tables["按月份"]) - 1)
        add_pie(writer.sheets["按区域"], "各区域销售额占比", len(tables["按区域"]) - 1)
        add_bar(writer.sheets["按销售员"], "销售员业绩对比", len(tables["按销售员"]) - 1)

    print("-" * 44)
    print(f"报表已生成：{OUT_FILE}")
    print(f"工作表（{len(tables)} 个）：{'、'.join(tables)}")
    print("图表（3 张）：月度柱状图、区域饼图、销售员柱状图")


if __name__ == "__main__":
    main()